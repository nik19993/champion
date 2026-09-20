#!/usr/bin/env python3
"""Build the Champion methodology catalogue from the approved DOCX sources."""

from __future__ import annotations

import hashlib
import json
import re
import uuid
from collections import Counter
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
TASKS_DOC = Path(
    "/workspace/scratch/c99be3f20f28/upload/Перелік тактичних завдань U6-U15.docx"
)
GLOSSARY_DOC = Path(
    "/workspace/scratch/c99be3f20f28/upload/Словник тактичних термінів футболу.docx"
)
DATA_DIR = ROOT / "dist" / "data"
SEED_DIR = ROOT / "supabase" / "seed"
NAMESPACE = uuid.UUID("4e535a3c-56f6-4fa2-88aa-0ca303ac217f")

EXPECTED_TASK_COUNTS = {
    6: {"attack": 20, "defense": 6},
    7: {"attack": 19, "defense": 10},
    8: {"attack": 23, "defense": 9},
    9: {"attack": 19, "defense": 17},
    10: {"attack": 17, "defense": 19},
    11: {"attack": 21, "defense": 15},
    12: {"attack": 19, "defense": 17},
    13: {"attack": 13, "defense": 19},
    14: {"attack": 12, "defense": 16},
    15: {"attack": 11, "defense": 16},
}

GLOSSARY_CATEGORIES = {
    "ЗАГАЛЬНІ ПОНЯТТЯ": "Загальні поняття",
    "ІНДИВІДУАЛЬНА ТАКТИКА НАПАДУ": "Індивідуальна тактика нападу",
    "ГРУПОВА ТАКТИКА НАПАДУ": "Групова тактика нападу",
    "КОМАНДНА ТАКТИКА НАПАДУ": "Командна тактика нападу",
    "СТАНДАРТНІ ПОЛОЖЕННЯ В НАПАДІ": "Стандартні положення в нападі",
    "ІНДИВІДУАЛЬНА ТАКТИКА ЗАХИСТУ": "Індивідуальна тактика захисту",
    "ГРУПОВА ТАКТИКА ЗАХИСТУ": "Групова тактика захисту",
    "КОМАНДНА ТАКТИКА ЗАХИСТУ": "Командна тактика захисту",
    "СТАНДАРТНІ ПОЛОЖЕННЯ В ЗАХИСТІ": "Стандартні положення в захисті",
}


def clean(value: str) -> str:
    return " ".join((value or "").replace("\u00a0", " ").split())


def source_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def parse_tasks(path: Path) -> list[dict]:
    document = Document(path)
    tasks: list[dict] = []
    age: int | None = None
    game_format = ""
    phase = ""
    section = ""
    paragraphs = document.paragraphs
    i = 0

    while i < len(paragraphs):
        paragraph = paragraphs[i]
        text = clean(paragraph.text)
        age_match = re.match(r"^U(\d{1,2})\s*•\s*([^•]+)\s*•", text)
        if age_match and paragraph.style.name == "Heading 1":
            age = int(age_match.group(1))
            game_format = clean(age_match.group(2))
            i += 1
            continue
        if paragraph.style.name == "Heading 2":
            if text.startswith("ТАКТИКА НАПАДУ"):
                phase = "attack"
            elif text.startswith("ТАКТИКА ЗАХИСТУ"):
                phase = "defense"
            i += 1
            continue
        if paragraph.style.name == "Heading 3":
            if text.startswith("Індивідуальна тактика"):
                section = "individual"
            elif text.startswith("Групова тактика"):
                section = "group"
            elif text.startswith("Командна тактика"):
                section = "team"
            i += 1
            continue

        task_match = re.match(r"^(\d+)\.\s+(.+)$", text)
        if task_match and age and phase and section:
            j = i + 1
            while j < len(paragraphs) and not clean(paragraphs[j].text):
                j += 1
            if j >= len(paragraphs):
                raise ValueError(f"Завдання без пояснення: {text}")
            instruction = clean(paragraphs[j].text)
            if not instruction.startswith("Навчити"):
                raise ValueError(f"Неочікуване пояснення після «{text}»: {instruction}")
            sequence = int(task_match.group(1))
            component_full = clean(task_match.group(2))
            component, _, focus = component_full.partition("→")
            key = f"curriculum-u{age:02d}-{phase}-{sequence:03d}"
            tasks.append(
                {
                    "id": key,
                    "dbId": str(uuid.uuid5(NAMESPACE, f"task:{key}")),
                    "age": age,
                    "phase": phase,
                    "section": section,
                    "component": component_full,
                    "baseComponent": clean(component),
                    "focus": clean(focus),
                    "title": instruction,
                    "sequence": sequence,
                    "gameFormat": game_format,
                    "source": "Перелік тактичних завдань U6–U15",
                }
            )
            i = j + 1
            continue
        i += 1

    counts = Counter((task["age"], task["phase"]) for task in tasks)
    expected = {
        (age_value, phase_value): count
        for age_value, phase_counts in EXPECTED_TASK_COUNTS.items()
        for phase_value, count in phase_counts.items()
    }
    if counts != Counter(expected):
        raise ValueError(f"Кількість завдань не збігається: {dict(counts)}")
    if len(tasks) != 318 or len({task["id"] for task in tasks}) != 318:
        raise ValueError("Очікувалося 318 унікальних навчальних завдань")
    return tasks


def iter_document_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def parse_glossary(path: Path) -> list[dict]:
    document = Document(path)
    category = ""
    entries: list[dict] = []

    for block in iter_document_blocks(document):
        if isinstance(block, Paragraph):
            heading = clean(block.text)
            match = re.match(r"^\d+\.\s+(.+)$", heading)
            if match:
                category = GLOSSARY_CATEGORIES.get(clean(match.group(1)).upper(), "")
            continue
        if not category or len(block.columns) < 2:
            continue
        for row in block.rows[1:]:
            term_parts = [clean(p.text) for p in row.cells[0].paragraphs if clean(p.text)]
            definition = clean(row.cells[1].text)
            if not term_parts or not definition:
                continue
            term = term_parts[0]
            english = " / ".join(term_parts[1:])
            key = f"glossary-{len(entries) + 1:03d}"
            entries.append(
                {
                    "id": key,
                    "dbId": str(
                        uuid.uuid5(
                            NAMESPACE, f"term:{category.casefold()}:{term.casefold()}"
                        )
                    ),
                    "term": term,
                    "english": english,
                    "definition": definition,
                    "category": category,
                    "scope": "shared",
                    "ownerId": None,
                    "source": "Словник тактичних термінів футболу",
                }
            )

    duplicate_keys = [
        key
        for key, count in Counter(
            (e["term"], e["category"]) for e in entries
        ).items()
        if count > 1
    ]
    category_counts = Counter(entry["category"] for entry in entries)
    if len(entries) != 110:
        raise ValueError(f"Очікувалося 110 термінів, отримано {len(entries)}")
    if duplicate_keys:
        raise ValueError(f"Повторювані терміни в одному розділі: {duplicate_keys}")
    if set(category_counts) != set(GLOSSARY_CATEGORIES.values()):
        raise ValueError(f"Неповний перелік розділів словника: {dict(category_counts)}")
    return entries


def build_sql(tasks: list[dict], entries: list[dict]) -> str:
    task_rows = [
        {
            "id": task["dbId"],
            "phase": "defence" if task["phase"] == "defense" else task["phase"],
            "level": task["section"],
            "component": task["component"],
            "title": task["title"],
            "description": f"{task['source']} · U{task['age']} · формат {task['gameFormat']}",
            "age_from": task["age"],
            "age_to": task["age"],
            "sort_order": task["age"] * 1000 + (500 if task["phase"] == "defense" else 0) + task["sequence"],
            "is_active": True,
        }
        for task in tasks
    ]
    term_rows = [
        {
            "id": entry["dbId"],
            "term": entry["term"],
            "definition": entry["definition"],
            "category": entry["category"],
            "source": f"{entry['source']}" + (f" · {entry['english']}" if entry["english"] else ""),
            "is_active": True,
        }
        for entry in entries
    ]
    tasks_json = json.dumps(task_rows, ensure_ascii=False, separators=(",", ":"))
    terms_json = json.dumps(term_rows, ensure_ascii=False, separators=(",", ":"))
    return f"""-- Generated from the approved Champion methodology documents.
-- Re-running this seed is safe: deterministic UUIDs and upserts prevent duplicates.
begin;

with source as (
  select *
  from jsonb_to_recordset($champion_tasks${tasks_json}$champion_tasks$::jsonb) as x(
    id uuid,
    phase public.methodology_phase,
    level public.methodology_level,
    component text,
    title text,
    description text,
    age_from smallint,
    age_to smallint,
    sort_order integer,
    is_active boolean
  )
)
insert into public.methodology_tasks
  (id, phase, level, component, title, description, age_from, age_to, sort_order, is_active)
select id, phase, level, component, title, description, age_from, age_to, sort_order, is_active
from source
on conflict (id) do update set
  phase = excluded.phase,
  level = excluded.level,
  component = excluded.component,
  title = excluded.title,
  description = excluded.description,
  age_from = excluded.age_from,
  age_to = excluded.age_to,
  sort_order = excluded.sort_order,
  is_active = excluded.is_active,
  updated_at = now();

with source as (
  select *
  from jsonb_to_recordset($champion_terms${terms_json}$champion_terms$::jsonb) as x(
    id uuid,
    term text,
    definition text,
    category text,
    source text,
    is_active boolean
  )
)
insert into public.terminology_entries
  (id, term, definition, category, source, is_active)
select id, term, definition, category, source, is_active
from source
on conflict (term, category) do update set
  definition = excluded.definition,
  category = excluded.category,
  source = excluded.source,
  is_active = excluded.is_active,
  updated_at = now();

commit;
"""


def write_outputs(tasks: list[dict], entries: list[dict]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    SEED_DIR.mkdir(parents=True, exist_ok=True)
    metadata = {
        "version": "2026-09-18",
        "taskCount": len(tasks),
        "termCount": len(entries),
        "taskSourceSha256": source_hash(TASKS_DOC),
        "glossarySourceSha256": source_hash(GLOSSARY_DOC),
        "taskCounts": {
            f"U{age}": EXPECTED_TASK_COUNTS[age] for age in sorted(EXPECTED_TASK_COUNTS)
        },
        "termCategoryCounts": dict(Counter(entry["category"] for entry in entries)),
    }
    payload = {"metadata": metadata, "tasks": tasks, "terms": entries}
    (DATA_DIR / "methodology-data.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (DATA_DIR / "methodology-data.js").write_text(
        "window.CHAMPION_METHODOLOGY = "
        + json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        + ";\n",
        encoding="utf-8",
    )
    (SEED_DIR / "20260918_methodology_content.sql").write_text(
        build_sql(tasks, entries), encoding="utf-8"
    )


def main() -> None:
    tasks = parse_tasks(TASKS_DOC)
    entries = parse_glossary(GLOSSARY_DOC)
    write_outputs(tasks, entries)
    print(f"Generated {len(tasks)} tasks and {len(entries)} terminology entries.")


if __name__ == "__main__":
    main()
